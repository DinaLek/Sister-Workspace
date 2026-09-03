from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(__file__).with_name('instagram-content-strategy-he.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.15
for name,size,before,after,color in [('Heading 1',20,20,6,'000000'),('Heading 2',16,18,6,'000000'),('Heading 3',14,16,4,'434343')]:
    s=styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=False; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Number']:
    s=styles[name]; s.font.name='Arial'; s.font.size=Pt(11); s.paragraph_format.left_indent=Inches(.5); s.paragraph_format.first_line_indent=Inches(-.25); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.15

def rtl(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    pPr=p._p.get_or_add_pPr(); bidi=OxmlElement('w:bidi'); bidi.set(qn('w:val'),'1'); pPr.append(bidi)

def title(text):
    p=doc.add_paragraph(); rtl(p); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(26); r.font.bold=False; r.font.color.rgb=RGBColor(0,0,0)

def p(text='', bold_prefix=None):
    q=doc.add_paragraph(); rtl(q)
    if bold_prefix and text.startswith(bold_prefix):
        q.add_run(bold_prefix).bold=True; q.add_run(text[len(bold_prefix):])
    else: q.add_run(text)
    return q

def h(text,level=1):
    q=doc.add_paragraph(text,style=f'Heading {level}'); rtl(q); return q

def bullet(text):
    q=doc.add_paragraph(text,style='List Bullet'); rtl(q); return q

def numbered(text):
    q=doc.add_paragraph(text,style='List Number'); rtl(q); return q

def set_cell(cell,text,bold=False):
    cell.text=''; q=cell.paragraphs[0]; rtl(q); r=q.add_run(str(text)); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(9.5); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    tcPr=cell._tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for side,val in [('top','80'),('bottom','80'),('start','120'),('end','120')]:
        el=OxmlElement('w:'+side); el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa'); mar.append(el)

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.RIGHT; t.autofit=False
    for i,x in enumerate(headers): set_cell(t.rows[0].cells[i],x,True)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): set_cell(cells[i],x)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    tblPr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'DADCE0'); borders.append(e)
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

title('אסטרטגיית תוכן ל־Instagram — 8.8.27')
p('ניתוח חשבונות השראה, אסטרטגיית צמיחה ותוכנית אימות ל־90 יום')
p('עודכן: 1 בספטמבר 2026 | שוק יעד: ארה״ב והעולם דובר האנגלית | חשבון faceless באנגלית')

h('תקציר מנהלים')
p('ההזדמנות אינה לבנות עוד עמוד מוטיבציה, אלא מותג תוכן שמושך אנשים דרך שאיפה להתפתחות אישית ומוביל אותם בהדרגה אל ביצוע בעולם אמיתי. עמודי המדיה הגדולים מוכיחים שתוכן קצר, רגשי וקל לשיתוף יכול לייצר הפצה רחבה. המותגים הרלוונטיים יותר מוכיחים שהמרה נוצרת כאשר אותה שאיפה מקבלת ritual, כלי או מערכת.')
p('המלצה מרכזית: למצב את החשבון בטריטוריה של Personal Growth That Survives Real Life. ההבטחה אינה להספיק יותר, אלא להמשיך לבנות את החיים החשובים לך גם כשהשבוע משתנה. בחודש הראשון יש להוביל עם תוכן רחב ושיתופי, אך לשמור לפחות 35% מהתוכן מחובר ל־small steps, realistic routines, life areas ו־adaptive progress כדי שלא יצטבר קהל לא רלוונטי.')
bullet('היעד הראשי: צמיחה אורגנית מהירה באמצעות שיתופים, שמירות וחשיפה ללא־עוקבים.')
bullet('היעד האסטרטגי: לזהות אילו מסרים מביאים קהל שגם מזדהה עם בעיית המוצר.')
bullet('הסיכון המרכזי: עוקבים רבים סביב חדשות, ציטוטים או מוטיבציה כללית ללא כוונת שימוש.')
bullet('היתרון התוכני: החיבור הייחודי בין becoming לבין realistic progress בשבוע לא מושלם.')

h('1. מתודולוגיה ומגבלות')
p('המחקר שילב נתונים בטבלת העבודה, פרופילים ציבוריים וניתוח מדגם של עד 30 פרסומים אחרונים בחמשת חשבונות העומק. בכל עשרת החשבונות נבדקו הביוגרפיה, גודל הקהל, קטגוריית התוכן, המוצר ומבנה המשפך. ניתוח העומק קודד נושא, פורמט, סוג הוק, רגש, CTA, מטרת התוכן וסימני ביצוע גלויים.')
bullet('חשבונות עומק: @mindset.therapy, @intelligentchange, @headway_app, @legacyacademyio ו־@bestselfunlocked.')
bullet('המספרים משקפים את הנתונים הגלויים ב־Instagram ביום הבדיקה; הם משתנים לאורך זמן.')
bullet('Instagram אינו מציג בכל חשבון נתוני reach, saves, shares או views היסטוריים מלאים. לכן ביצועים שאינם גלויים מסומנים כאינדיקציה מבנית, לא כהוכחה כמותית.')
bullet('המחקר מחקה עקרונות ומנגנונים בלבד. אין להעתיק ניסוח, עיצוב או חומר מוגן.')

h('2. מפת עשרת חשבונות ההשראה')
table(['חשבון','קהל נצפה','מודל תוכן','מנגנון מסחרי','לקח מרכזי'],[
['@mindset.therapy','11.3M','מדיה ויראלית: ציטוטים, חדשות, תרבות וספורט','שותפויות ותנועה ל־Viralspot','היקף ותזמון מייצרים reach, אך התאמת המוצר נמוכה אם מחקים אקטואליה.'],
['@intelligentchange','337K','שאלות, affirmations, relationships, gratitude וממים','יומנים, אפליקציות ומוצרים פיזיים','מחבר זהות ורגש ל־ritual של חמש דקות ולמוצר.'],
['@dailystoic','3.4M','ציטוטים, חכמה מעשית וקטעי מומחה','ספרים, newsletter, קורסים ואירועים','רעיון פילוסופי עקבי מאפשר אקו־סיסטם מסחרי רחב.'],
['@headway_app','4.5M','קרוסלות פסיכולוגיה, ממים, שאלות ותובנות ספרים','אפליקציית microlearning','המוצר נטמע בתוכן כדרך להעמקה, לא כפרסומת נפרדת.'],
['@bestselfco','145K','intentional living, journaling, focus וקשרים','יומנים, planners, decks וכלי focus','הקשר החזק ביותר בין becoming, זמן מכוון וכלי פיזי.'],
['@legacyacademyio','275K','Reels מוטיבציוניים faceless עם סיפורים ומטאפורות','אפליקציית self-improvement','מכונת תוכן פשוטה עם CTA מבוסס מילת קוד ו־DM.'],
['@mindjournal','91.7K','רפלקציה, גבריות, בריאות נפשית ו־journaling','יומן פיזי לגברים','נישה ברורה יוצרת התאמה גבוהה, אך מגבילה קהל.'],
['@growthday','68K','ציטוטי מומחים, ביצועים והרגלים','membership ואפליקציה','authority דרך מומחים; פחות מתאים למנגנון faceless זול.'],
['@bestselfunlocked','49.5K','soft structure, realistic routines והרגלים','משאבים דיגיטליים ו־affiliate','המסר הקרוב ביותר לבעיה, אך הפעילות אינה עקבית.'],
['@sheisdedicatedtogrowth','79','תזכורות growth לנשים ו־glow-up','affiliate ליומנים ו־planners','קל לייצור אך גנרי; גודל החשבון אינו מוכיח הצלחה.']
],[1.15,.65,1.55,1.45,1.7])

h('3. ניתוח עומק של חמישה מודלים')
h('@mindset.therapy — מנוע מדיה, לא מודל מותג מלא',2)
p('המדגם מציג קצב גבוה מאוד, שילוב של חדשות בידור וספורט, ציטוטי ידוענים, עובדות וקרוסלות “swipe for more”. כוחו של החשבון הוא packaging: כותרת גדולה, מתח מיידי, דמות מוכרת והבטחה למידע נוסף. התוכן הרלוונטי למוצר מופיע רק כחלק קטן מהפיד.')
bullet('הוק חוזר: חדשות, אמירה מפתיעה, דמות מפורסמת או “What nobody tells you…”.')
bullet('רגש: סקרנות, הפתעה, נוסטלגיה והזדהות מהירה.')
bullet('CTA: swipe, share או קריאת caption; מעט מאוד מעבר למוצר.')
bullet('לאמץ: קריאות מיידית, רעיון אחד בכל פוסט, packaging חד וקצב ניסויים גבוה.')
bullet('להימנע: אקטואליה שאינה קשורה להבטחת המותג; היא יכולה לייצר reach אך לזהם את הקהל.')

h('@intelligentchange — השראה שמקבלת ritual ומוצר',2)
p('החשבון נע בין affirmations קצרים, שאלות לשיחה, prompts לרפלקציה, תוכן מערכות יחסים וממים. ההמרה אינה קפיצה חדה: הטקסט מתחיל בזהות או כאב, ואז מציג journaling של חמש דקות כהמשך טבעי. קמפיינים משתמשים במילת קוד בתגובות כדי לייצר intent ו־DM.')
bullet('הוק חוזר: “Repeat after me”, מספר שאלות, משפט becoming או סיטואציה רגשית מוכרת.')
bullet('רגש: תקווה, אינטימיות, ביטחון עצמי ושייכות.')
bullet('CTA: comment keyword, link in bio, download או shop.')
bullet('לאמץ: חיבור בין זהות יומיומית לפעולה קטנה; סדרות שאלות; מעבר רך מתוכן למוצר.')
bullet('להתאים: במקום journaling, להוביל מ־reflection לצעד קטן שמקבל מקום בשבוע.')

h('@headway_app — מותג אפליקציה שמתנהג כמותג תרבות',2)
p('Headway אינו מפרסם רק סיכומי ספרים. המדגם כולל relationships, parenting, perfectionism, rest, people pleasing, burnout וממים. הוא משתמש בשפה אינטרנטית, שאלות והומור כדי להרחיב את הקהל, ואז מציע ספר או את האפליקציה כדרך להעמיק.')
bullet('הוק חוזר: “5 questions…”, “The trap / The way out”, תיוג עצמי ומם מזדהה.')
bullet('פורמטים: קרוסלות של 5–8 שקופיות, Reels קצרים, ממים ותוכן אינטראקטיבי.')
bullet('CTA: תגובה, בחירה, תיוג חבר, או המלצת ספר/אפליקציה.')
bullet('לאמץ: שפה לא־פרסומית, מגוון נושאים תחת umbrella אחד, והפיכת המוצר ל־next step.')
bullet('סיכון: רוחב נושאים גדול מחייב הבטחת מותג חזקה כדי לא להפוך למגזין כללי.')

h('@legacyacademyio — מכונת Reels והמרת תגובות ל־DM',2)
p('זהו המודל התפעולי הפשוט ביותר: Reels faceless המבוססים על קטע חזותי, מטאפורה או סיטואציה, עם טקסט מוטיבציוני קצר. ה־caption מרחיב את הלקח ומסתיים במילת קוד כגון ACTION, FOCUS או CONSISTENT. החיבור למוצר ישיר יותר מבחשבונות מדיה.')
bullet('הוק חוזר: שאלה מאתגרת, ניגוד, מטאפורה חזותית או אמת לא נעימה.')
bullet('רגש: דחיפות, נחישות, פחד מהחמצה והעצמה.')
bullet('CTA: comment keyword לקבלת כלי; follow for more.')
bullet('לאמץ: תבנית הפקה חוזרת, CTA מדיד וחיבור בין intent ל־action.')
bullet('להימנע: שפה שיפוטית של discipline בלבד; היא מנוגדת להבטחת גמישות וחמלה בתקופות קשות.')

h('@bestselfunlocked — התאמת כאב גבוהה, הוכחת הפצה חלשה',2)
p('התוכן עוסק בהרגלים קטנים, progress over perfection, soft structure, עומס ומערכות שמתאימות לחיים אמיתיים. זהו הקול הקרוב ביותר למוצר, אך חלק גדול מהמדגם ישן והחשבון אינו מדגים קצב צמיחה עדכני. הוא מקור למסרים, לא benchmark לביצועים.')
bullet('הוק חוזר: “If you keep starting and quitting…”, “Here’s why”, “Save this”.')
bullet('רגש: הקלה, הבנה עצמית, מסוגלות ורוגע.')
bullet('CTA: שמירה, תגובה עם מילת קוד וקבלת tracker חינמי.')
bullet('לאמץ: small enough to do, life happens, system that fits you, ו־minimum progress.')
bullet('לשפר: קופי קצר יותר, packaging חד יותר ורלוונטיות ניטרלית מגדרית.')

h('4. מה מייצר חשיפה לעומת התאמה למוצר')
table(['מנגנון','פוטנציאל חשיפה','התאמה למוצר','החלטה'],[
['ציטוט אוניברסלי קצר','גבוה','נמוכה–בינונית','להשתמש רק כשיש זווית של progress או real life.'],
['חדשות וידוענים','גבוה מאוד','נמוכה','לא להפוך לעמוד חדשות; להשתמש רק ב־RTM רלוונטי נדיר.'],
['מם על perfectionism או hustle','גבוה','גבוהה','פורמט ליבה: shareable ומסנן לקהל הנכון.'],
['קרוסלת שאלות / checklist','בינוני–גבוה','גבוהה','פורמט ליבה לשמירות ולבניית סמכות.'],
['מטאפורת Reel faceless','גבוה','בינונית–גבוהה','להשתמש עם מסקנה ייחודית של adaptive progress.'],
['הדגמת מוצר ישירה','נמוך בתחילת הדרך','גבוהה','5% בלבד; להרחיב לאחר שיש מוצר ברור ועדויות.'],
['Comment keyword','בינוני','גבוהה','לבדוק לאחר יצירת lead magnet או waitlist.']
],[1.6,1.05,1.15,2.7])

h('5. קהל התוכן והמיצוב')
h('קהל ראשי',2)
p('Intentional Life Builders: בני ובנות 27–45, לרוב אנשי מקצוע, עצמאים, יזמים או עובדים בעלי אוטונומיה יחסית בזמן. הם צורכים ספרים, פודקאסטים ותוכן התפתחות אישית; משתמשים ביומן; ורוצים להתקדם במקביל בבריאות, למידה, קשרים, קריירה ופרויקטים אישיים. הבעיה שלהם אינה חוסר שאיפה אלא הפער בין כוונה לחיים בפועל.')
h('קהל משני',2)
p('אנשים שאוהבים מבנה חזותי אך חווים קיבולת משתנה, כולל חלק מבעלי ADHD ו־working parents. אין למצב את המוצר כפתרון רפואי או כמיועד ל־ADHD לפני מחקר ייעודי.')
h('Anti-audience',2)
p('צרכני מוטיבציה פסיביים שאינם רוצים מערכת ביצוע; אנשים שנלחצים מעצם השימוש בלו״ז; וקהל productivity שמטרתו היחידה לדחוס יותר עבודה.')
h('ניסוחי מיצוב מומלצים באנגלית',2)
bullet('Account promise: “Build the life that matters — even when the week doesn’t go to plan.”')
bullet('Content proposition: “Practical personal growth for busy, unpredictable lives.”')
bullet('Positioning line: “Where becoming meets real life.”')
bullet('Alternative test: “Progress that adapts when life does.”')
bullet('Bio draft: “Personal growth for real life. Small steps, balanced weeks, visible progress. Build what matters — even when plans change.”')

h('6. חמשת עמודי התוכן')
table(['עמוד תוכן','כאב/שאיפה','פורמטים והוקים','רגש ו־CTA','חיבור עתידי למוצר'],[
['Progress Over Perfect','שבוע שהשתבש מרגיש כמו כישלון','Reels מטאפוריים; quotes; “You are not behind…”','הקלה ותקווה; share/save','Minimum, tracking והתקדמות מצטברת.'],
['Make Room for What Matters','עבודה דוחקת בריאות, קשרים ולמידה','קרוסלות life areas; “If it matters, it needs a place”','בהירות; לבחור תחום/לתייג','קטגוריות חיים ותעדוף.'],
['Soft Structure, Strong Follow-Through','מבנה נוקשה נשבר; חוסר מבנה יוצר כאוס','checklists, memes, before/after; “Your plan needs room to move”','מסוגלות; save/comment','קיבולת, Target/Minimum ותכנון שבועי.'],
['Tiny Actions, Compounding Change','שאיפות גדולות אינן הופכות לפעולות','micro-challenges, habit menus; “Make it small enough…”','מומנטום; try today/save','אבני דרך וצעדים שבועיים.'],
['Life Happened. Now What?','בלת״ם אחד מפיל את כל היום','scenario carousels; reschedule demos; “The plan changed. The goal didn’t.”','ביטחון; choose A/B','adaptive rescheduling והצעת next best step.']
],[1.25,1.35,1.65,1.15,1.1])

h('7. תמהיל, פורמטים וסדרות')
p('תמהיל פתיחה: 60% תוכן שיתופי לצמיחה, 25% תוכן שימושי לשמירה, 10% תוכן problem/solution ו־5% תוכן מוצר או waitlist. בפועל, בשבוע של 10 פרסומים: 6 shareable, שניים עד שלושה saveable, אחד problem/solution ופוסט מוצר אחד בכל שבועיים.')
h('סדרות קבועות',2)
table(['שם הסדרה','פורמט','דוגמת הוק באנגלית','CTA'],[
['When Life Happens','Reel טקסטואלי 7–12 שניות','“The plan changed. The goal didn’t.”','Send this to someone rebuilding their week.'],
['The Minimum Still Counts','קרוסלה 5–6 שקופיות','“What progress looks like on a low-capacity day.”','Save this for a hard week.'],
['Make Room','קרוסלת checklist','“Five things you say matter — but never make the calendar.”','Which one needs a place this week?'],
['Not Lazy. Overloaded.','מם / single graphic','“You don’t need more discipline. You need a plan that knows your capacity.”','Share if this is the week you’re having.'],
['Tiny Step Menu','קרוסלה לפי תחומי חיים','“Choose one 10-minute way to move your life forward today.”','Comment the number you’re choosing.'],
['Weekly Reset','קרוסלה ritual','“Before you plan the week, decide what ‘enough’ means.”','Save for Sunday.'],
['Future You, Scheduled','quote + caption','“Your future self is built in ordinary Tuesdays.”','Follow for growth that fits real life.'],
['Plan B Is Still a Plan','Reel/קרוסלת תרחיש','“Missed the workout? Here are three ways the goal can survive.”','Which Plan B would you choose?']
],[1.2,1.25,2.75,1.3])

h('8. כללי קריאייטיב וקופי')
bullet('כל פוסט מבטא רעיון אחד שניתן להבין בתוך שנייה אחת מהפריים הראשון.')
bullet('שפה קצרה, ישירה, אמפתית ולא טיפולית. לא להשתמש בהבטחות רפואיות או באבחון.')
bullet('להעדיף ניגוד: perfect plan / real life, motivation / system, all-or-nothing / minimum.')
bullet('לא להשתמש בבושה, “no excuses” או האשמת הקהל בחוסר discipline.')
bullet('Reels: 7–15 שניות, טקסט גדול, 2–4 משפטים, loop טבעי וויזואל פשוט עם זכויות שימוש.')
bullet('קרוסלות: 5–7 שקופיות; cover חד, 3–5 נקודות, סיום עם פעולה אחת.')
bullet('Captions: 60–140 מילים ברוב הפוסטים; שורת פתיחה עצמאית; CTA יחיד.')
bullet('עיצוב: טיפוגרפיה נקייה, ניגודיות גבוהה, צבעוניות רגועה אך לא wellness גנרית. להימנע מחיקוי trade dress של המתחרים.')

h('9. מטריצת בחירה עתידית לסוכן')
p('כל רעיון תוכן עתידי צריך להיווצר משילוב של תא אחד מכל שכבה. המבנה נועד לשמש בשלב הבא כ־input schema לסוכן השבועי.')
table(['שכבה','אפשרויות'],[
['Audience moment','Sunday planning; bad day; missed task; overwhelmed morning; end-of-week reflection; new goal; unexpected interruption'],
['Pillar','Progress Over Perfect; Make Room; Soft Structure; Tiny Actions; Life Happened'],
['Core tension','perfect vs possible; motivation vs system; busy vs meaningful; rigid vs adaptive; intention vs calendar'],
['Format','short Reel; 5-slide carousel; meme; quote; checklist; scenario A/B'],
['Hook pattern','contrarian truth; identity statement; relatable confession; numbered list; question; before/after'],
['Emotion','relief; hope; recognition; urgency; confidence; curiosity'],
['CTA','share; save; comment choice; follow; join waitlist; comment keyword'],
['Product bridge','life categories; capacity; minimum; rescheduling; progress graph; weekly reset']
],[1.5,5.0])

h('10. דוגמאות תוכן באנגלית')
h('Shareable examples',2)
bullet('“A disrupted week is not a wasted week.”')
bullet('“You don’t need to do everything. You need to keep something important alive.”')
bullet('“The goal was never a perfect routine. The goal was a life that keeps moving.”')
bullet('“If your plan only works on good days, it isn’t a plan yet.”')
bullet('“Small progress in five parts of your life can matter more than one perfect workday.”')
h('Carousel examples',2)
numbered('Cover: “What progress looks like when you have 20% capacity.”')
numbered('Health: a ten-minute walk still protects the identity you are building.')
numbered('Learning: two pages still keep the thread alive.')
numbered('Relationships: one thoughtful message still counts as showing up.')
numbered('Growth: minimum is not failure; it is continuity.')
numbered('CTA: “Save this for the week that refuses to cooperate.”')
h('Product-bridge examples',2)
bullet('“Your calendar knows when you are free. It doesn’t know what kind of life you are trying to build.”')
bullet('“Most planners track what you finished. What if yours protected what matters before the week began?”')
bullet('“When life moves the plan, the next step should move with it.”')

h('11. תוכנית אימות ל־90 יום')
table(['תקופה','שאלת המחקר','ניסויים','החלטה בסוף התקופה'],[
['ימים 1–30','איזה שילוב מייצר reach וקהל נכון?','שלוש טריטוריות מסר; Reels מול carousels; שני סגנונות hook','לבחור 3 סדרות מנצחות ולבטל פורמטים חלשים.'],
['ימים 31–60','איזה כאב יוצר intent?','soft structure מול balanced life; minimum מול adaptive planning; CTA save מול comment','לבחור מסר ליבה ו־lead magnet או waitlist angle.'],
['ימים 61–90','האם הקהל מגיב להצעת המוצר?','problem/solution, concept demo, weekly reset ו־waitlist','להחליט אם Instagram הוא acquisition channel או בעיקר awareness.']
],[.95,2.0,2.15,1.4])
h('לוח ניסויים שבועי',2)
bullet('שבוע 1: אותו נושא, Hook של הקלה מול Hook ניגודי.')
bullet('שבוע 2: personal growth כללי מול balanced progress.')
bullet('שבוע 3: quote/Reel מול קרוסלת checklist.')
bullet('שבוע 4: progress over perfect מול adaptive planning.')
bullet('שבועות 5–8: לחזור רק על שני הצירים המובילים ולבדוק CTA אחד בכל פעם.')
bullet('שבועות 9–12: להוסיף מסר מוצר, concept demo ורשימת המתנה בהדרגה.')

h('12. מדדים וכללי החלטה')
table(['מדד','למה הוא חשוב','כלל החלטה ראשוני'],[
['Non-follower reach','מודד הפצה מעבר לקהל הקיים','להשוות לחציון 10 הפוסטים האחרונים, לא למספר מוחלט.'],
['Shares / reach','הסיגנל המרכזי לתוכן זהותי ו־viral','פי 1.5 מהחציון = מועמד לסדרה חוזרת.'],
['Saves / reach','מודד שימושיות וכוונת חזרה','פי 1.5 מהחציון = להרחיב את הנושא/הפורמט.'],
['Profile visits / reach','מודד סקרנות כלפי ההבטחה','עלייה עקבית מצביעה על positioning חד.'],
['Follows / profile visits','מודד התאמה בין הפוסט לחשבון','נמוך = פער בין ויראליות להבטחת הפרופיל.'],
['Comments with intent','מודד כאב, בחירה או רצון בכלי','לקודד איכותנית, לא רק לספור.'],
['Waitlist clicks/signups','מודד התאמה מסחרית','להתחיל למדוד גם כשהכמות קטנה.']
],[1.35,2.35,2.8])
p('סדרה מנצחת: לפחות שלושה פוסטים באותה סדרה, כששניים מהם עוברים את חציון החשבון ב־50% במדד הראשי. סדרה בינונית: reach סביר אך ללא shares/saves; משנים Hook או packaging לפני שמבטלים. סדרה חלשה: שלושה ניסיונות מתחת לחציון גם לאחר שינוי Hook — עוצרים לחודש.')

h('13. המלצת פרסום לחודש הראשון')
bullet('10 פוסטים בשבוע: 5 Reels faceless, 3 קרוסלות, 2 גרפיקות/ממים.')
bullet('Stories: 3–5 ימים בשבוע, בעיקר polls, A/B ושאלות; לא נדרש רצף יומי מלא בתחילת הדרך.')
bullet('שני נושאי growth רחבים, שני נושאי realistic routines, שני progress-over-perfect, שני practical/saveable, אחד scenario adaptive ואחד post ניסוי.')
bullet('פוסט מוצר ישיר אחד בכל שבועיים בלבד, עד שיש prototype שניתן להדגים או waitlist ברורה.')
bullet('לא לפרסם אוטומטית ללא סקירת איכות. בשלב הבא הסוכן יכין חבילה שבועית לאישור מרוכז.')

h('14. החלטות אסטרטגיות')
numbered('לא לבנות חשבון productivity. לבנות מותג personal-growth execution.')
numbered('לא להעתיק את רוחב התוכן של Mindset Therapy; לאמץ רק את חדות ה־packaging.')
numbered('להשתמש ב־Legacy Academy כמודל הפקה, ב־Intelligent Change כמודל המרה וב־Headway כמודל גיוון.')
numbered('להשתמש ב־BestSelfUnlocked כמקור לשפת הכאב, לא כהוכחת צמיחה.')
numbered('להחזיק בכל שבוע תוכן רחב לצמיחה לצד תוכן שמסנן לקהל שאוהב מבנה והתקדמות.')
numbered('לבנות את הסוכן רק לאחר חודש ראשון של נתונים, או לבנות אותו עם מנגנון ניסויים ולא עם “נוסחה מנצחת” שטרם הוכחה.')

h('מקורות')
for src in [
'Google Sheet — חשבונות השראה ושיווק וניתוח אסטרטגי: https://docs.google.com/spreadsheets/d/1HuH_YJUzb0ZB54wd2_qtVqPMszXBWrTMEehHrx7RhTU/edit',
'Instagram profiles reviewed: instagram.com/mindset.therapy, /intelligentchange, /dailystoic, /headway_app, /bestselfco, /legacyacademyio, /mindjournal, /growthday, /bestselfunlocked, /sheisdedicatedtogrowth',
'מחקר השוק והמיצוב הקודם של 8.8.27, Sister Marketing, אוגוסט 2026.'
]: bullet(src)

doc.save(OUT)
lines=[]
for block in doc.element.body:
    tag=block.tag.rsplit('}',1)[-1]
    if tag=='p':
        text=''.join(node.text or '' for node in block.iter() if node.tag.rsplit('}',1)[-1]=='t').strip()
        if text: lines.append(text)
    elif tag=='tbl':
        for row in block.findall('.//w:tr', {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
            cells=[]
            for cell in row.findall('./w:tc', {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
                cells.append(''.join(node.text or '' for node in cell.iter() if node.tag.rsplit('}',1)[-1]=='t').strip())
            if cells: lines.append(' | '.join(cells))
Path(__file__).with_name('instagram-content-strategy-he.txt').write_text('\n'.join(lines),encoding='utf-8')
print(OUT)
